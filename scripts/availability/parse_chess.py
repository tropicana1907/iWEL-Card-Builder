"""Шахматка (xlsx из Битрикса) → наличие квартир без персональных данных.

Использование: python parse_chess.py <project> <file.xlsx> [...] > out.json
В результат попадают только номер, корпус/блок, этаж, площадь и статус.
ФИО, телефоны и примечания из шахматки НЕ выводятся.
"""
import json, re, sys
import openpyxl

# Голубой 00B0F0 в обеих шахматках = первый этаж / «не продаётся» (в легенде Империала, блок 3)
FIXED_COLORS = {'00B0F0': 'not_for_sale'}

LEGEND_WORDS = {'купленная': 'sold', 'продана': 'sold', 'бронь': 'reserved', 'не продается': 'not_for_sale',
                'рассрочка': 'sold'}


def color(c):
    f = c.fill
    if not f or f.fill_type is None:
        return '-'
    fg = f.fgColor
    if fg.type == 'theme':
        return f'th{fg.theme}/{round(fg.tint, 2)}'
    if fg.type == 'indexed':
        return f'i{fg.indexed}'
    return fg.rgb[2:]


def num(v):
    if v is None:
        return None
    m = re.search(r'\d+([.,]\d+)?', str(v))
    return float(m.group().replace(',', '.')) if m else None


def floor_of(v):
    m = re.fullmatch(r'\s*(\d+)\s*(эт)?\s*', str(v or ''))
    return int(m.group(1)) if m else None


def legend(ws):
    out = {}
    for row in ws.iter_rows(min_row=1, max_row=8):
        for c in row:
            if isinstance(c.value, str):
                w = c.value.strip().lower()
                for k, st in LEGEND_WORDS.items():
                    if w.startswith(k) and c.column > 1:
                        out[color(ws.cell(row=c.row, column=c.column - 1))] = st
    return out


def sections(ws):
    """Находит таблицы шахматки: ячейка с названием корпуса/блока, под ней строка площадей и строки этажей."""
    found = []
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and re.match(r'^\s*(корпус|блок)\s*\S+', c.value.strip(), re.I):
                found.append(c)
    return found


def parse_sheet(ws, project):
    leg = {**FIXED_COLORS, **legend(ws)}
    res, skipped = [], []
    for title in sections(ws):
        name = re.sub(r'\s+', ' ', title.value.strip())
        c0, r0 = title.column, title.row
        # строка площадей: первая строка ниже, где правее заголовка есть числа
        area_row = next((r for r in range(r0 + 1, r0 + 5)
                         if sum(1 for c in range(c0 + 1, c0 + 14) if num(ws.cell(row=r, column=c).value)) >= 3), None)
        if not area_row:
            continue
        cols = []
        for c in range(c0 + 1, c0 + 16):
            a = num(ws.cell(row=area_row, column=c).value)
            if a:
                cols.append((c, a))
            elif cols and ws.cell(row=area_row, column=c).value in (None, '\xa0'):
                break
        # список владельцев: «кв №» в той же колонке ниже
        owners, list_row = {}, None
        for r in range(area_row, ws.max_row + 1):
            v = ws.cell(row=r, column=c0).value
            if isinstance(v, str) and v.strip().lower().startswith('кв'):
                list_row = r
                break
        if list_row:
            for r in range(list_row + 1, ws.max_row + 1):
                k = ws.cell(row=r, column=c0).value
                if k is None:
                    continue
                nm = ws.cell(row=r, column=c0 + 1).value
                owners[re.sub(r'\D', '', str(k).split('\\')[0]) + ('/' + str(k).split('\\')[1] if '\\' in str(k) else '')] = bool(nm and str(nm).strip())
        sec = []
        for r in range(area_row + 1, (list_row or ws.max_row + 1)):
            fl = floor_of(ws.cell(row=r, column=c0).value)
            if fl is None:
                continue
            for c, a in cols:
                v = ws.cell(row=r, column=c).value
                s = str(v or '').strip().replace('\xa0', '')
                if not s or 'омм' in s.lower():
                    continue
                parts = re.findall(r'\d+', s)
                if not parts:
                    continue
                apt = parts[0] + ('/' + parts[1] if '\\' in s or '/' in s else '')
                col = color(ws.cell(row=r, column=c))
                st = leg.get(col)
                if st is None:
                    st = 'sold' if owners.get(apt) else 'free'
                elif st == 'free' and owners.get(apt):
                    st = 'sold'
                if st in ('sold',) and not owners.get(apt) and leg.get(col) is None:
                    st = 'free'
                sec.append({'section': name, 'floor': fl, 'apt': apt, 'area': a, 'status': st})
        # Заготовка шахматки (на каждом этаже «Кв1…Кв7») — корпус ещё не в продаже
        nums = [x['apt'] for x in sec]
        if nums and len(set(nums)) < 0.7 * len(nums):
            skipped.append(name)
            continue
        res += sec
    return res, skipped


def main():
    project, files = sys.argv[1], sys.argv[2:]
    apts, skipped = [], []
    for f in files:
        wb = openpyxl.load_workbook(f)
        for ws in wb.worksheets:
            a, s = parse_sheet(ws, project)
            apts += a
            skipped += s
    print(json.dumps({'project': project, 'apartments': apts, 'skipped': skipped}, ensure_ascii=False))


if __name__ == '__main__':
    main()
